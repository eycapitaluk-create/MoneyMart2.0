#!/usr/bin/env python3
"""Numbers 와이드 종가 시트 → import-us-close-wide-xlsx.mjs 가 기대하는 형식의 xlsx 로 내보내기."""
import argparse
import sys
from pathlib import Path

from numbers_parser import Document
from openpyxl import Workbook


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("numbers_path", type=Path)
    ap.add_argument(
        "-o",
        "--out",
        type=Path,
        required=True,
        help="출력 .xlsx 경로",
    )
    ap.add_argument(
        "--sheet-index",
        type=int,
        default=0,
        help="문서 내 시트 인덱스 (기본 0 = 첫 시트)",
    )
    args = ap.parse_args()

    doc = Document(str(args.numbers_path))
    sheet = doc.sheets[args.sheet_index]
    if not sheet.tables:
        print("No tables in sheet", file=sys.stderr)
        sys.exit(1)
    table = sheet.tables[0]

    wb = Workbook()
    ws = wb.active
    ws.title = sheet.name[:31] or "Close Prices"

    for ri in range(table.num_rows):
        row = [table.cell(ri, ci).value for ci in range(table.num_cols)]
        if ri <= 2:
            # Ticker, Date (Text), SYM… → Ticker, SYM… (날짜 열 제거)
            ws.append([row[0]] + row[2:])
        else:
            # datetime, text date, prices… → 첫 열만 날짜로 유지
            ws.append([row[0]] + row[2:])

    wb.save(str(args.out))
    print(f"Wrote {args.out} ({table.num_rows} rows × {table.num_cols} cols → wide)")


if __name__ == "__main__":
    main()
